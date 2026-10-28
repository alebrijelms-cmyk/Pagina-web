"""
Proyecto 5 - Distribución y monitoreo de pozos de riego agrícola (REPDA/CONAGUA)
Ordenamiento aleatorio + búsqueda lineal/binaria con benchmarking en nanosegundos.

Uso:  python proyecto_pozos.py [Pozos_Riego_15000.xlsx]
"""
import json
import math
import os
import random
import sys
import time
from collections import Counter

from openpyxl import load_workbook

ARCHIVO = sys.argv[1] if len(sys.argv) > 1 else "Pozos_Riego_15000.xlsx"
CONTADOR = "contador_metodos.json"
CAMPOS = ["id_titulo_concesion", "volumen_extraido_m3", "fecha_concesion", "nombre_titular", "estado",
          "municipio", "profundidad_pozo_m", "estatus_pozo", "caudal_litros_segundo", "uso_principal"]
CLAVES = {1: "id_titulo_concesion", 2: "volumen_extraido_m3", 3: "fecha_concesion"}
METODOS = {  # métodos asignados según el campo clave (Sección 3 del documento)
    1: ["Radix Sort", "Counting Sort"],
    2: ["Bucket Sort"],
    3: ["Quick Sort", "Merge Sort", "Heap Sort"],
}


class Cont:
    """Contador de comparaciones (c) y movimientos/intercambios (m)."""
    def __init__(self):
        self.c = 0
        self.m = 0


# ───────────────────────── Algoritmos de ordenamiento ─────────────────────────
# Todos reciben `filas` (lista de listas) y el índice `k` de la columna clave; devuelven lista ordenada.

def counting_sort(filas, k, ct):
    mx = max(f[k] for f in filas)
    cuenta = [0] * (mx + 1)
    for f in filas:
        cuenta[f[k]] += 1
        ct.m += 1
    for i in range(1, mx + 1):
        cuenta[i] += cuenta[i - 1]
    salida = [None] * len(filas)
    for f in reversed(filas):
        cuenta[f[k]] -= 1
        salida[cuenta[f[k]]] = f
        ct.m += 1
    return salida


def radix_sort(filas, k, ct):
    mx, exp = max(f[k] for f in filas), 1
    while mx // exp > 0:
        cuenta = [0] * 10
        for f in filas:
            cuenta[(f[k] // exp) % 10] += 1
        for i in range(1, 10):
            cuenta[i] += cuenta[i - 1]
        salida = [None] * len(filas)
        for f in reversed(filas):
            d = (f[k] // exp) % 10
            cuenta[d] -= 1
            salida[cuenta[d]] = f
            ct.m += 1
        filas, exp = salida, exp * 10
    return filas


def bucket_sort(filas, k, ct):
    n = len(filas)
    nb = max(1, int(math.sqrt(n)) * 8)
    mx = max(f[k] for f in filas)
    esc = math.log1p(mx) or 1  # escala logarítmica: el volumen está muy sesgado
    cubetas = [[] for _ in range(nb)]
    for f in filas:
        cubetas[min(nb - 1, int(math.log1p(f[k]) / esc * (nb - 1)))].append(f)
        ct.m += 1
    salida = []
    for cb in cubetas:  # insertion sort dentro de cada cubeta
        for i in range(1, len(cb)):
            x, j = cb[i], i - 1
            while j >= 0:
                ct.c += 1
                if cb[j][k] > x[k]:
                    cb[j + 1] = cb[j]
                    ct.m += 1
                    j -= 1
                else:
                    break
            cb[j + 1] = x
            ct.m += 1
        salida.extend(cb)
    return salida


def quick_sort(filas, k, ct):
    a = filas[:]
    pila = [(0, len(a) - 1)]
    while pila:
        lo, hi = pila.pop()
        while lo < hi:
            mid = (lo + hi) // 2  # mediana de tres
            ct.c += 3
            if a[mid][k] < a[lo][k]:
                a[lo], a[mid] = a[mid], a[lo]; ct.m += 1
            if a[hi][k] < a[lo][k]:
                a[lo], a[hi] = a[hi], a[lo]; ct.m += 1
            if a[hi][k] < a[mid][k]:
                a[mid], a[hi] = a[hi], a[mid]; ct.m += 1
            p = a[mid][k]
            i, j = lo - 1, hi + 1
            while True:  # partición de Hoare (maneja bien duplicados)
                i += 1
                ct.c += 1
                while a[i][k] < p:
                    i += 1; ct.c += 1
                j -= 1
                ct.c += 1
                while a[j][k] > p:
                    j -= 1; ct.c += 1
                if i >= j:
                    break
                a[i], a[j] = a[j], a[i]; ct.m += 1
            if j - lo < hi - j - 1:       # se apila el lado grande, se itera el pequeño
                pila.append((j + 1, hi)); hi = j
            else:
                pila.append((lo, j)); lo = j + 1
    return a


def merge_sort(filas, k, ct):
    if len(filas) <= 1:
        return filas
    med = len(filas) // 2
    izq, der = merge_sort(filas[:med], k, ct), merge_sort(filas[med:], k, ct)
    res, i, j = [], 0, 0
    while i < len(izq) and j < len(der):
        ct.c += 1
        if izq[i][k] <= der[j][k]:
            res.append(izq[i]); i += 1
        else:
            res.append(der[j]); j += 1
        ct.m += 1
    resto = izq[i:] + der[j:]
    ct.m += len(resto)
    return res + resto


def heap_sort(filas, k, ct):
    a, n = filas[:], len(filas)

    def sift(i, tam):
        while True:
            mayor, l, r = i, 2 * i + 1, 2 * i + 2
            if l < tam:
                ct.c += 1
                if a[l][k] > a[mayor][k]:
                    mayor = l
            if r < tam:
                ct.c += 1
                if a[r][k] > a[mayor][k]:
                    mayor = r
            if mayor == i:
                return
            a[i], a[mayor] = a[mayor], a[i]; ct.m += 1
            i = mayor

    for i in range(n // 2 - 1, -1, -1):
        sift(i, n)
    for fin in range(n - 1, 0, -1):
        a[0], a[fin] = a[fin], a[0]; ct.m += 1
        sift(0, fin)
    return a


ALGORITMOS = {"Radix Sort": radix_sort, "Counting Sort": counting_sort, "Bucket Sort": bucket_sort,
              "Quick Sort": quick_sort, "Merge Sort": merge_sort, "Heap Sort": heap_sort}


# ───────────────────────────── Búsquedas ─────────────────────────────

def busqueda_lineal(filas, idx, valor):
    ct, hallados = Cont(), []
    for f in filas:
        ct.c += 1
        if f[idx] == valor:
            hallados.append(f)
    return hallados, ct


def busqueda_binaria(filas, idx, valor):
    ct, lo, hi = Cont(), 0, len(filas) - 1
    while lo <= hi:
        mid = (lo + hi) // 2
        ct.c += 1
        if filas[mid][idx] == valor:
            a = b = mid  # se extiende a duplicados vecinos
            while a > 0 and filas[a - 1][idx] == valor: a -= 1; ct.c += 1
            while b < len(filas) - 1 and filas[b + 1][idx] == valor: b += 1; ct.c += 1
            return filas[a:b + 1], ct
        ct.c += 1
        if filas[mid][idx] < valor:
            lo = mid + 1
        else:
            hi = mid - 1
    return [], ct


# ───────────────────────────── Utilidades ─────────────────────────────

def cargar(ruta):
    wb = load_workbook(ruta, read_only=True)
    filas = [list(r) for r in wb["Datos"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb.close()
    return filas


def tabla(filas, titulo):
    print(f"\n── {titulo} ──")
    anchos = [10, 10, 9, 30, 16, 20, 8, 13, 8, 9]
    cab = ["id_titulo", "vol_m3", "fecha", "titular", "estado", "municipio", "prof_m", "estatus", "L/s", "uso"]
    print(" ".join(c[:w].ljust(w) for c, w in zip(cab, anchos)))
    print(" ".join("-" * w for w in anchos))
    for f in filas:
        print(" ".join(str(v)[:w].ljust(w) for v, w in zip(f, anchos)))


def convertir(texto, campo):
    if campo in ("id_titulo_concesion", "volumen_extraido_m3", "fecha_concesion"):
        return int(texto)
    if campo in ("profundidad_pozo_m", "caudal_litros_segundo"):
        return float(texto)
    return texto.strip()


def leer(msg):
    try:
        return input(msg)
    except EOFError:
        return ""


def main():
    print("=" * 78)
    print(" PROYECTO 5: DISTRIBUCIÓN Y MONITOREO DE POZOS DE RIEGO AGRÍCOLA Y CONCESIONES")
    print(" Fuente: REPDA - CONAGUA | 15,000 títulos | Ordenamiento y búsqueda con benchmarking")
    print("=" * 78)

    t0 = time.perf_counter_ns()
    filas = cargar(ARCHIVO)
    t_carga = time.perf_counter_ns() - t0
    print(f"\nSe cargaron {len(filas):,} registros en memoria RAM en {t_carga:,} nanosegundos.")
    tabla(filas[:30], "Primeros 30 registros")
    tabla(filas[-30:], "Últimos 30 registros")

    # ── Disparador aleatorio de ordenación
    k_sel = random.randint(1, 3)
    campo = CLAVES[k_sel]
    metodo = random.choice(METODOS[k_sel])
    idx = CAMPOS.index(campo)
    ct = Cont()
    t0 = time.perf_counter_ns()
    ordenado = ALGORITMOS[metodo](filas, idx, ct)
    t_ord = time.perf_counter_ns() - t0
    assert all(ordenado[i][idx] <= ordenado[i + 1][idx] for i in range(len(ordenado) - 1)), "Orden incorrecto"
    print(f"\nCampo clave sorteado: {k_sel} | Método sorteado: {metodo}")
    print(f"Se ordenaron {len(ordenado):,} registros por el campo {campo} usando el método {metodo} "
          f"en {t_ord:,} nanosegundos. Comparaciones: {ct.c:,}, Intercambios/Movimientos: {ct.m:,}")
    print(f"→ Para el Excel de registro: Ejecución | {k_sel} | {campo} | {metodo} | {t_ord} ns")

    # ── Disparador aleatorio de búsqueda
    tipo = random.choice(["Lineal", "Binaria"])
    print(f"\nTipo de búsqueda sorteada: {tipo}")
    if tipo == "Binaria":
        print(f"Se usará el campo clave ya ordenado: {campo}")
        base, idx_b, campo_b = ordenado, idx, campo
    else:
        print("Campos disponibles:", ", ".join(CAMPOS))
        campo_b = ""
        while campo_b not in CAMPOS:
            campo_b = leer("Campo donde buscar: ").strip() or random.choice(CAMPOS)
        base, idx_b = filas, CAMPOS.index(campo_b)
    valor = None
    while valor is None:
        txt = leer(f"Dato a buscar en {campo_b}: ")
        if txt == "":  # sin entrada interactiva: toma un dato existente para demostrar
            txt = str(random.choice(base)[idx_b]); print(f"(sin entrada, se usa: {txt})")
        try:
            valor = convertir(txt, campo_b)
        except ValueError:
            print("Valor no válido para ese campo.")
    t0 = time.perf_counter_ns()
    hallados, cts = (busqueda_binaria if tipo == "Binaria" else busqueda_lineal)(base, idx_b, valor)
    t_bus = time.perf_counter_ns() - t0
    print(f"\nBúsqueda {tipo} en {campo_b} = {valor}: {len(hallados)} coincidencia(s), "
          f"{cts.c:,} comparaciones, {t_bus:,} nanosegundos.")
    if hallados:
        tabla(hallados[:10], "Resultados (máx. 10)")

    # ── Método Alfa (el más seleccionado), acumulado entre ejecuciones
    cont = Counter(json.load(open(CONTADOR))) if os.path.exists(CONTADOR) else Counter()
    cont[metodo] += 1
    json.dump(cont, open(CONTADOR, "w"), indent=2)
    print("\n── Selección de métodos de ordenamiento (acumulado) ──")
    for m, n in cont.most_common():
        print(f"  {m:<14} {n} vez/veces")
    alfa = cont.most_common(1)[0]
    print(f"\n★ Método Alfa: {alfa[0]} ({alfa[1]} selección/es)")


if __name__ == "__main__":
    main()
